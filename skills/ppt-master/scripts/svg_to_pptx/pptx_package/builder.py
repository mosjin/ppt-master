"""Core PPTX assembly: create_pptx_with_native_svg."""

from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import posixpath
import re
import shutil
import stat
import subprocess
import tempfile
import uuid
import zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.util import Emu

from ..drawingml.converter import convert_svg_to_slide_shapes
from .dimensions import (
    CANVAS_FORMATS,
    get_slide_dimensions, get_pixel_dimensions,
    get_viewbox_dimensions, detect_format_from_svg,
)
from .media import (
    PNG_RENDERER,
    get_png_renderer_info, convert_svg_to_png, convert_svg_to_png_cached,
)
from .notes import (
    markdown_to_plain_text,
    create_notes_master_rels_xml,
    create_notes_master_xml,
    create_notes_slide_xml,
    create_notes_slide_rels_xml,
)
from .narration import (
    AUDIO_CONTENT_TYPES,
    AUDIO_REL_TYPE,
    AUDIO_MARKER_PNG_BYTES,
    IMAGE_REL_TYPE,
    MEDIA_REL_TYPE,
    apply_recorded_timing,
    inject_narration,
    next_shape_id,
    probe_audio_duration,
)
from .slide_xml import (
    ANIMATIONS_AVAILABLE, TRANSITIONS,
    create_slide_xml_with_svg, create_slide_rels_xml,
    link_shape_xml,
)
from .svg_link_extractor import extract_links

_HYPERLINK_REL_TYPE = (
    'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
)

# Re-import create_transition_xml only if available
try:
    from pptx_animations import (
        create_transition_xml,
        create_sequence_timing_xml,
        pick_animation_effect,
    )
except ImportError:
    create_transition_xml = None
    create_sequence_timing_xml = None
    pick_animation_effect = None


SLIDE_LAYOUT_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
)
SLIDE_MASTER_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
)
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
DML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"

for _prefix, _uri in (("p", PML_NS), ("a", DML_NS), ("r", REL_NS), ("p14", P14_NS)):
    try:
        ET.register_namespace(_prefix, _uri)
    except (ValueError, AttributeError):
        pass


@dataclass(frozen=True)
class PptxStructureContext:
    """Resolved base package structure reused when slide XML is regenerated."""

    slide_layout_targets: dict[int, str]
    slide_master_parts: dict[int, str]

    def slide_layout_target(self, slide_num: int) -> str:
        """Return the slide layout target for a generated slide."""
        try:
            return self.slide_layout_targets[slide_num]
        except KeyError as exc:
            raise RuntimeError(
                f"Missing slide layout relationship for generated slide {slide_num}"
            ) from exc

    def slide_master_part(self, slide_num: int) -> str:
        """Return the slide master package part for a generated slide."""
        try:
            return self.slide_master_parts[slide_num]
        except KeyError as exc:
            raise RuntimeError(
                f"Missing slide master relationship for generated slide {slide_num}"
            ) from exc


def _relationship_attrs(elem: ET.Element) -> dict[str, str]:
    return {key.rsplit("}", 1)[-1]: value for key, value in elem.attrib.items()}


def _resolve_package_target(source_part: str, target: str) -> str:
    """Resolve a relationship target relative to a package part path."""
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _relationships_path_for_part(extract_dir: Path, part_name: str) -> Path:
    """Return the package relationship sidecar path for a part name."""
    path = Path(part_name)
    return extract_dir / path.parent / "_rels" / f"{path.name}.rels"


def _find_relationship_target(
    rels_path: Path,
    rel_type: str,
) -> str | None:
    """Find the first relationship target for a relationship type."""
    if not rels_path.exists():
        return None
    root = ET.parse(rels_path).getroot()
    for elem in root:
        attrs = _relationship_attrs(elem)
        if attrs.get("Type") == rel_type:
            return attrs.get("Target")
    return None


def _read_relationships(rels_path: Path) -> dict[str, dict[str, str]]:
    """Return relationship attributes keyed by rId."""
    if not rels_path.exists():
        return {}
    root = ET.parse(rels_path).getroot()
    rels: dict[str, dict[str, str]] = {}
    for elem in root:
        attrs = _relationship_attrs(elem)
        rel_id = attrs.get("Id")
        if rel_id:
            rels[rel_id] = attrs
    return rels


def _find_relationship_id(
    rels_path: Path,
    rel_type: str,
    target: str,
) -> str | None:
    """Find an existing relationship by type and target."""
    for rel_id, attrs in _read_relationships(rels_path).items():
        if attrs.get("Type") == rel_type and attrs.get("Target") == target:
            return rel_id
    return None


def _read_slide_layout_targets(extract_dir: Path, slide_count: int) -> PptxStructureContext:
    """Read the actual layout relationship target for every generated slide."""
    slide_layout_targets: dict[int, str] = {}
    slide_master_parts: dict[int, str] = {}
    rels_dir = extract_dir / "ppt" / "slides" / "_rels"
    for slide_num in range(1, slide_count + 1):
        rels_path = rels_dir / f"slide{slide_num}.xml.rels"
        if not rels_path.exists():
            raise RuntimeError(f"Missing slide relationship file: {rels_path}")
        target = _find_relationship_target(rels_path, SLIDE_LAYOUT_REL_TYPE)
        if not target:
            raise RuntimeError(f"Slide {slide_num} has no slide layout relationship")
        slide_layout_targets[slide_num] = target

        slide_part = f"ppt/slides/slide{slide_num}.xml"
        layout_part = _resolve_package_target(slide_part, target)
        layout_rels_path = _relationships_path_for_part(extract_dir, layout_part)
        master_target = _find_relationship_target(layout_rels_path, SLIDE_MASTER_REL_TYPE)
        if not master_target:
            raise RuntimeError(
                f"Slide {slide_num} layout has no slide master relationship"
            )
        slide_master_parts[slide_num] = _resolve_package_target(layout_part, master_target)
    return PptxStructureContext(
        slide_layout_targets=slide_layout_targets,
        slide_master_parts=slide_master_parts,
    )


_SLIDE_BACKGROUND_RE = re.compile(
    r"(?P<prefix><p:cSld\b[^>]*>\s*)"
    r"(?P<bg><p:bg\b.*?</p:bg>)"
    r"(?P<suffix>\s*<p:spTree\b)",
    re.DOTALL,
)


def _extract_slide_background_xml(slide_xml: str) -> str | None:
    """Return the slide-level p:bg XML when it directly precedes spTree."""
    match = _SLIDE_BACKGROUND_RE.search(slide_xml)
    return match.group("bg") if match else None


def _remove_slide_background_xml(slide_xml: str) -> str:
    """Remove a promoted slide-level p:bg from cSld."""
    return _SLIDE_BACKGROUND_RE.sub(r"\g<prefix>\g<suffix>", slide_xml, count=1)


def _put_background_on_master(master_xml: str, background_xml: str) -> str | None:
    """Replace or insert the master-level p:bg before the master spTree.

    Returns None when the master carries a p:bg the canonical pattern cannot
    replace; inserting there would leave two p:bg children under p:cSld.
    """
    match = _SLIDE_BACKGROUND_RE.search(master_xml)
    if match:
        return (
            master_xml[:match.start("bg")]
            + background_xml
            + master_xml[match.end("bg"):]
        )
    if "<p:bg" in master_xml:
        return None

    cslide_match = re.search(r"(<p:cSld\b[^>]*>)", master_xml)
    if not cslide_match:
        raise RuntimeError("Slide master XML has no p:cSld element")
    return (
        master_xml[:cslide_match.end()]
        + background_xml
        + master_xml[cslide_match.end():]
    )


def _promote_common_slide_backgrounds_to_masters(
    extract_dir: Path,
    structure: PptxStructureContext,
    slide_count: int,
    *,
    verbose: bool = False,
) -> int:
    """Promote identical slide backgrounds to their shared slide master."""
    slides_by_master: dict[str, list[int]] = {}
    for slide_num in range(1, slide_count + 1):
        master_part = structure.slide_master_part(slide_num)
        slides_by_master.setdefault(master_part, []).append(slide_num)

    promoted = 0
    for master_part, slide_nums in slides_by_master.items():
        slide_backgrounds: dict[int, str] = {}
        for slide_num in slide_nums:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            slide_xml = slide_path.read_text(encoding="utf-8")
            background_xml = _extract_slide_background_xml(slide_xml)
            if not background_xml:
                slide_backgrounds = {}
                break
            slide_backgrounds[slide_num] = background_xml

        if not slide_backgrounds:
            continue
        unique_backgrounds = set(slide_backgrounds.values())
        if len(unique_backgrounds) != 1:
            continue

        background_xml = next(iter(unique_backgrounds))
        master_path = extract_dir / master_part
        master_xml = master_path.read_text(encoding="utf-8")
        promoted_master_xml = _put_background_on_master(master_xml, background_xml)
        if promoted_master_xml is None:
            continue
        master_path.write_text(promoted_master_xml, encoding="utf-8")

        for slide_num in slide_nums:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            slide_xml = slide_path.read_text(encoding="utf-8")
            slide_path.write_text(
                _remove_slide_background_xml(slide_xml),
                encoding="utf-8",
            )
            promoted += 1

    if verbose and promoted:
        print(f"  Baseline master background: promoted {promoted} slide background(s)")
    return promoted


_CHROME_TRACE_TOKENS = (
    "logo",
    "footer",
    "header",
    "watermark",
    "chrome",
    "pagenumber",
    "slidenumber",
    "pagenum",
    "slidenum",
)
_TOP_LEVEL_SHAPE_TAGS = {
    f"{{{PML_NS}}}sp",
    f"{{{PML_NS}}}grpSp",
    f"{{{PML_NS}}}pic",
    f"{{{PML_NS}}}cxnSp",
    f"{{{PML_NS}}}graphicFrame",
}
_REL_ATTRS = {
    f"{{{REL_NS}}}embed",
    f"{{{REL_NS}}}link",
    f"{{{REL_NS}}}id",
}


def _chrome_token_from_svg_id(svg_id: str | None) -> str | None:
    """Return the baseline chrome token encoded in a source SVG id."""
    if not svg_id:
        return None
    lower = svg_id.lower()
    compact = re.sub(r"[-_\s]+", "", lower)
    if compact in _CHROME_TRACE_TOKENS:
        return compact
    split_tokens = {token for token in re.split(r"[-_\s]+", lower) if token}
    for token in _CHROME_TRACE_TOKENS:
        if token in split_tokens:
            return token
    return None


def _trace_chrome_shape_ids(
    trace: dict[str, Any] | None,
) -> dict[str, list[str]]:
    """Map chrome token to generated top-level shape ids for one slide."""
    result: dict[str, list[str]] = {}
    if not trace:
        return result
    for event in trace.get("events", []):
        if event.get("decision") != "native":
            continue
        token = _chrome_token_from_svg_id(event.get("id"))
        shape_id = event.get("shape_id")
        if token and shape_id is not None:
            shape_ids = result.setdefault(token, [])
            normalized_shape_id = str(shape_id)
            if normalized_shape_id not in shape_ids:
                shape_ids.append(normalized_shape_id)
    return result


def _shape_id(elem: ET.Element) -> str | None:
    for cnv in elem.iter(f"{{{PML_NS}}}cNvPr"):
        return cnv.attrib.get("id")
    return None


def _top_level_shapes_by_id(root: ET.Element) -> dict[str, ET.Element]:
    sp_tree = root.find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
    if sp_tree is None:
        return {}
    shapes: dict[str, ET.Element] = {}
    for child in list(sp_tree):
        if child.tag not in _TOP_LEVEL_SHAPE_TAGS:
            continue
        shape_id = _shape_id(child)
        if shape_id:
            shapes[shape_id] = child
    return shapes


def _timing_shape_ids(root: ET.Element) -> set[str]:
    """Return slide-local shape ids referenced by animation timing."""
    return {
        elem.attrib["spid"]
        for elem in root.findall(f".//{{{PML_NS}}}timing//{{{PML_NS}}}spTgt")
        if elem.attrib.get("spid")
    }


def _relationship_ids_in_shape(elem: ET.Element) -> set[str]:
    rel_ids: set[str] = set()
    for node in elem.iter():
        for attr_name, value in node.attrib.items():
            if attr_name in _REL_ATTRS and value:
                rel_ids.add(value)
    return rel_ids


def _shape_relationships_supported(
    elem: ET.Element,
    rels: dict[str, dict[str, str]],
) -> bool:
    """Only image relationships are safe to copy into a slide master here."""
    for rel_id in _relationship_ids_in_shape(elem):
        attrs = rels.get(rel_id)
        if not attrs:
            return False
        if attrs.get("TargetMode"):
            return False
        if attrs.get("Type") != IMAGE_REL_TYPE:
            return False
    return True


def _canonical_shape_xml(
    elem: ET.Element,
    rels: dict[str, dict[str, str]],
) -> bytes:
    """Canonicalize ids and relationship ids for cross-slide equality."""
    clone = ET.fromstring(ET.tostring(elem, encoding="utf-8"))
    for cnv in clone.iter(f"{{{PML_NS}}}cNvPr"):
        cnv.set("id", "ID")
        # Generated names include the slide-local shape id (for example,
        # ``Image 2`` versus ``Image 8``) but do not affect rendering.
        if "name" in cnv.attrib:
            cnv.set("name", "NAME")
    for node in clone.iter():
        for attr_name, value in list(node.attrib.items()):
            if attr_name not in _REL_ATTRS:
                continue
            attrs = rels.get(value, {})
            node.set(
                attr_name,
                f"{attrs.get('Type', '')}|{attrs.get('Target', '')}",
            )
    return ET.tostring(clone, encoding="utf-8")


def _ensure_relationship(
    rels_path: Path,
    rel_type: str,
    target: str,
) -> str:
    existing = _find_relationship_id(rels_path, rel_type, target)
    if existing:
        return existing
    return _append_relationship(rels_path, rel_type, target)


def _copy_shape_relationships_to_master(
    elem: ET.Element,
    slide_rels: dict[str, dict[str, str]],
    master_rels_path: Path,
) -> ET.Element:
    """Clone a shape and retarget supported relationship ids to the master."""
    clone = ET.fromstring(ET.tostring(elem, encoding="utf-8"))
    for node in clone.iter():
        for attr_name, value in list(node.attrib.items()):
            if attr_name not in _REL_ATTRS:
                continue
            rel = slide_rels.get(value)
            if not rel:
                raise RuntimeError(f"Missing slide relationship for {value}")
            new_rid = _ensure_relationship(
                master_rels_path,
                rel["Type"],
                rel["Target"],
            )
            node.set(attr_name, new_rid)
    return clone


def _next_master_shape_id(master_xml: str) -> int:
    ids = [
        int(match)
        for match in re.findall(r"<p:cNvPr\b[^>]*\bid=\"(\d+)\"", master_xml)
    ]
    return max(ids, default=1) + 1


def _renumber_shape_ids(elem: ET.Element, start_id: int) -> None:
    next_id = start_id
    for cnv in elem.iter(f"{{{PML_NS}}}cNvPr"):
        cnv.set("id", str(next_id))
        next_id += 1


def _append_shape_to_master(master_path: Path, elem: ET.Element) -> None:
    master_xml = master_path.read_text(encoding="utf-8")
    _renumber_shape_ids(elem, _next_master_shape_id(master_xml))
    shape_xml = ET.tostring(elem, encoding="unicode")
    if "</p:spTree>" not in master_xml:
        raise RuntimeError(f"Slide master has no p:spTree: {master_path}")
    master_path.write_text(
        master_xml.replace("</p:spTree>", f"{shape_xml}\n</p:spTree>", 1),
        encoding="utf-8",
    )


def _write_xml_tree(path: Path, tree: ET.ElementTree) -> None:
    tree.write(path, encoding="utf-8", xml_declaration=True)


def _promote_common_chrome_shapes_to_masters(
    extract_dir: Path,
    structure: PptxStructureContext,
    slide_count: int,
    conversion_traces: list[dict[str, Any]] | None,
    *,
    verbose: bool = False,
) -> int:
    """Promote explicit repeated chrome SVG ids to their shared master."""
    if not conversion_traces:
        return 0
    trace_by_slide = {
        int(trace.get("slide_num", 0)): trace
        for trace in conversion_traces
        if trace.get("slide_num") is not None
    }
    if len(trace_by_slide) < slide_count:
        return 0

    slides_by_master: dict[str, list[int]] = {}
    for slide_num in range(1, slide_count + 1):
        master_part = structure.slide_master_part(slide_num)
        slides_by_master.setdefault(master_part, []).append(slide_num)

    promoted = 0
    promoted_roles = 0
    for master_part, slide_nums in slides_by_master.items():
        if len(slide_nums) < 2:
            continue
        slide_state: dict[int, dict[str, Any]] = {}
        for slide_num in slide_nums:
            slide_path = extract_dir / "ppt" / "slides" / f"slide{slide_num}.xml"
            rels_path = extract_dir / "ppt" / "slides" / "_rels" / f"slide{slide_num}.xml.rels"
            tree = ET.parse(slide_path)
            root = tree.getroot()
            slide_state[slide_num] = {
                "path": slide_path,
                "rels": _read_relationships(rels_path),
                "root": root,
                "shapes": _top_level_shapes_by_id(root),
                "timing_shape_ids": _timing_shape_ids(root),
                "tokens": _trace_chrome_shape_ids(trace_by_slide.get(slide_num)),
                "tree": tree,
            }

        promotions: list[tuple[str, dict[int, str]]] = []
        claimed_shape_ids: dict[int, set[str]] = {
            slide_num: set() for slide_num in slide_nums
        }
        all_tokens = sorted({
            token
            for state in slide_state.values()
            for token in state["tokens"]
        })
        for token in all_tokens:
            shape_ids_by_slide: dict[int, str] = {}
            canonical_by_slide: dict[int, bytes] = {}
            for slide_num in slide_nums:
                state = slide_state[slide_num]
                shape_ids = state["tokens"].get(token, [])
                if len(shape_ids) != 1:
                    shape_ids_by_slide = {}
                    break
                shape_id = shape_ids[0]
                shape = state["shapes"].get(shape_id)
                if shape is None:
                    shape_ids_by_slide = {}
                    break
                if shape_id in state["timing_shape_ids"]:
                    shape_ids_by_slide = {}
                    break
                if not _shape_relationships_supported(shape, state["rels"]):
                    shape_ids_by_slide = {}
                    break
                shape_ids_by_slide[slide_num] = shape_id
                canonical_by_slide[slide_num] = _canonical_shape_xml(
                    shape,
                    state["rels"],
                )
            if set(shape_ids_by_slide) != set(slide_nums):
                continue
            if len(set(canonical_by_slide.values())) != 1:
                continue
            # A flattened nested chrome group can emit several semantic trace
            # ids for the same generated DrawingML shape. Claim it once.
            if any(
                shape_ids_by_slide[slide_num] in claimed_shape_ids[slide_num]
                for slide_num in slide_nums
            ):
                continue
            for slide_num, shape_id in shape_ids_by_slide.items():
                claimed_shape_ids[slide_num].add(shape_id)
            promotions.append((token, shape_ids_by_slide))

        if not promotions:
            continue

        # Master shapes always render behind slide-local shapes. Preserve the
        # original z-order by promoting only a common leading chrome prefix;
        # overlay headers/footers remain slide-local.
        token_by_shape_id = {
            slide_num: {
                shape_ids[slide_num]: token
                for token, shape_ids in promotions
            }
            for slide_num in slide_nums
        }
        leading_token_orders: list[list[str]] = []
        for slide_num in slide_nums:
            order: list[str] = []
            for shape_id in slide_state[slide_num]["shapes"]:
                token = token_by_shape_id[slide_num].get(shape_id)
                if token is None:
                    break
                order.append(token)
            leading_token_orders.append(order)

        safe_tokens = list(leading_token_orders[0])
        for order in leading_token_orders[1:]:
            common_length = 0
            for expected, actual in zip(safe_tokens, order):
                if expected != actual:
                    break
                common_length += 1
            safe_tokens = safe_tokens[:common_length]
            if not safe_tokens:
                break
        promotion_by_token = {token: shape_ids for token, shape_ids in promotions}
        promotions = [
            (token, promotion_by_token[token])
            for token in safe_tokens
        ]

        if not promotions:
            continue

        master_path = extract_dir / master_part
        master_rels_path = _relationships_path_for_part(extract_dir, master_part)
        for _token, shape_ids_by_slide in promotions:
            first_slide = slide_nums[0]
            first_state = slide_state[first_slide]
            shape = first_state["shapes"][shape_ids_by_slide[first_slide]]
            master_shape = _copy_shape_relationships_to_master(
                shape,
                first_state["rels"],
                master_rels_path,
            )
            _append_shape_to_master(master_path, master_shape)
            promoted_roles += 1

            for slide_num, shape_id in shape_ids_by_slide.items():
                state = slide_state[slide_num]
                shape_to_remove = state["shapes"].get(shape_id)
                sp_tree = state["root"].find(f".//{{{PML_NS}}}cSld/{{{PML_NS}}}spTree")
                if sp_tree is not None and shape_to_remove is not None:
                    sp_tree.remove(shape_to_remove)
                    promoted += 1

        for state in slide_state.values():
            _write_xml_tree(state["path"], state["tree"])

    if verbose and promoted:
        print(
            "  Baseline master chrome: "
            f"promoted {promoted} slide shape(s) across {promoted_roles} shared object(s)"
        )
    return promoted


def _append_relationship(
    rels_path: Path,
    rel_type: str,
    target: str,
) -> str:
    """Append a relationship entry with the next available rId."""
    with open(rels_path, 'r', encoding='utf-8') as f:
        rels_content = f.read()

    rid_numbers = [int(match) for match in re.findall(r'Id="rId(\d+)"', rels_content)]
    next_rid = f'rId{max(rid_numbers, default=0) + 1}'
    rel_xml = (
        f'  <Relationship Id="{next_rid}" '
        f'Type="{rel_type}" Target="{target}"/>'
    )
    rels_content = rels_content.replace(
        '</Relationships>', rel_xml + '\n</Relationships>',
    )

    with open(rels_path, 'w', encoding='utf-8') as f:
        f.write(rels_content)

    return next_rid


def _add_default_content_type(content_types: str, extension: str, content_type: str) -> str:
    """Add a Default content type if it is not already present."""
    ext = extension.lstrip(".")
    if f'Extension="{ext}"' in content_types:
        return content_types
    entry = f'  <Default Extension="{ext}" ContentType="{content_type}"/>'
    override_pos = content_types.find('<Override ')
    if override_pos >= 0:
        return content_types[:override_pos] + entry + '\n' + content_types[override_pos:]
    return content_types.replace('</Types>', entry + '\n</Types>')


def _add_content_type_override(content_types: str, part_name: str, content_type: str) -> str:
    """Add an Override content type if it is not already present."""
    normalized = '/' + part_name.lstrip('/')
    if f'PartName="{normalized}"' in content_types:
        return content_types
    entry = f'  <Override PartName="{normalized}" ContentType="{content_type}"/>'
    return content_types.replace('</Types>', entry + '\n</Types>')


_IMAGE_CONTENT_TYPES = {
    'png': 'image/png',
    'jpg': 'image/jpeg',
    'jpeg': 'image/jpeg',
    'gif': 'image/gif',
    'webp': 'image/webp',
    'svg': 'image/svg+xml',
    'bmp': 'image/bmp',
    'emf': 'image/x-emf',
    'tif': 'image/tiff',
    'tiff': 'image/tiff',
    'wmf': 'image/x-wmf',
}


def _content_type_for_extension(ext: str) -> str:
    clean = ext.lower().lstrip('.')
    content_type = _IMAGE_CONTENT_TYPES.get(clean) or mimetypes.guess_type(f'x.{clean}')[0]
    if not content_type:
        raise ValueError(f"Unknown media content type for extension: {ext}")
    return content_type


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _create_writable_work_dir(output_path: Path) -> Path:
    """Create a real writable work directory for PPTX assembly."""
    parents = [output_path.parent, Path.cwd(), Path(tempfile.gettempdir())]
    seen: set[str] = set()
    errors: list[str] = []

    for parent in parents:
        parent = parent if str(parent) else Path(".")
        try:
            key = str(parent.resolve())
        except OSError:
            key = str(parent.absolute())
        if key in seen:
            continue
        seen.add(key)

        try:
            parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            errors.append(f"{parent}: cannot create parent ({exc})")
            continue

        for _ in range(3):
            work_dir = parent / f".pptx-build-{os.getpid()}-{uuid.uuid4().hex}"
            try:
                work_dir.mkdir(mode=0o700)
                probe_path = work_dir / ".write-probe"
                probe_path.write_text("ok", encoding="utf-8")
                probe_path.unlink()
                return work_dir
            except OSError as exc:
                errors.append(f"{work_dir}: {exc}")
                shutil.rmtree(work_dir, ignore_errors=True)

    details = "\n  - ".join(errors) if errors else "no candidate directories available"
    raise PermissionError(
        "Unable to create a writable PPTX work directory. "
        "Set the output path to a writable project directory or adjust sandbox permissions. "
        f"Tried:\n  - {details}"
    )


def _relax_output_permissions(output_path: Path) -> list[str]:
    """Make exported files readable outside the sandbox owner where possible."""
    warnings: list[str] = []

    try:
        current_mode = output_path.stat().st_mode
        readable_mode = (
            current_mode
            | stat.S_IRUSR
            | stat.S_IWUSR
            | stat.S_IRGRP
            | stat.S_IROTH
        )
        os.chmod(output_path, readable_mode)
    except OSError as exc:
        warnings.append(f"chmod skipped for {output_path}: {exc}")

    if os.name != 'nt':
        return warnings

    # Windows ACLs can remain sandbox-only even when the file mode looks sane.
    # Grant the built-in Users SID read access; the SID avoids localization
    # issues on non-English Windows installations.
    try:
        result = subprocess.run(
            ['icacls', str(output_path), '/grant', '*S-1-5-32-545:R'],
            capture_output=True,
            text=True,
            check=False,
        )
    except OSError as exc:
        warnings.append(f"icacls skipped for {output_path}: {exc}")
    else:
        if result.returncode != 0:
            message = (result.stderr or result.stdout or '').strip()
            details = f": {message}" if message else ''
            warnings.append(f"icacls failed for {output_path}{details}")

    return warnings


_NOTES_MASTER_REL_TYPE = (
    'http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesMaster'
)


def _ensure_notes_master(extract_dir: Path) -> None:
    """Create notesMaster parts and wire them into the presentation package."""
    ppt_dir = extract_dir / 'ppt'
    notes_masters_dir = ppt_dir / 'notesMasters'
    notes_masters_dir.mkdir(exist_ok=True)

    notes_master_path = notes_masters_dir / 'notesMaster1.xml'
    if not notes_master_path.exists():
        notes_master_path.write_text(create_notes_master_xml(), encoding='utf-8')

    theme_dir = ppt_dir / 'theme'
    theme_dir.mkdir(exist_ok=True)
    theme1_path = theme_dir / 'theme1.xml'
    theme2_path = theme_dir / 'theme2.xml'
    if not theme2_path.exists():
        if theme1_path.exists():
            shutil.copy2(theme1_path, theme2_path)
        else:
            raise RuntimeError('Cannot create notes theme: ppt/theme/theme1.xml is missing')

    notes_master_rels_dir = notes_masters_dir / '_rels'
    notes_master_rels_dir.mkdir(exist_ok=True)
    notes_master_rels_path = notes_master_rels_dir / 'notesMaster1.xml.rels'
    if not notes_master_rels_path.exists():
        notes_master_rels_path.write_text(
            create_notes_master_rels_xml(),
            encoding='utf-8',
        )

    presentation_rels_path = ppt_dir / '_rels' / 'presentation.xml.rels'
    notes_master_rid = _find_relationship_id(
        presentation_rels_path,
        _NOTES_MASTER_REL_TYPE,
        'notesMasters/notesMaster1.xml',
    )
    if notes_master_rid is None:
        notes_master_rid = _append_relationship(
            presentation_rels_path,
            _NOTES_MASTER_REL_TYPE,
            'notesMasters/notesMaster1.xml',
        )

    presentation_path = ppt_dir / 'presentation.xml'
    presentation_xml = presentation_path.read_text(encoding='utf-8')
    if '<p:notesMasterIdLst>' in presentation_xml:
        return
    notes_master_lst = (
        f'<p:notesMasterIdLst><p:notesMasterId r:id="{notes_master_rid}"/>'
        '</p:notesMasterIdLst>'
    )
    if '</p:sldMasterIdLst>' not in presentation_xml:
        raise RuntimeError('presentation.xml is missing p:sldMasterIdLst')
    presentation_xml = presentation_xml.replace(
        '</p:sldMasterIdLst>',
        '</p:sldMasterIdLst>' + notes_master_lst,
        1,
    )
    presentation_path.write_text(presentation_xml, encoding='utf-8')


def _to_float(value: Any, default: float) -> float:
    if value is None:
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        return default
    return number if number >= 0 else default


def _slide_config(animation_config: dict[str, Any] | None, svg_stem: str) -> dict[str, Any]:
    if not animation_config:
        return {}
    slides = _as_dict(animation_config.get('slides'))
    return _as_dict(slides.get(svg_stem))


def _slide_transition_settings(
    slide_cfg: dict[str, Any],
    transition: str | None,
    duration: float,
    auto_advance: float | None,
    cli_overrides: dict[str, bool],
) -> tuple[str | None, float, float | None]:
    trans_cfg = _as_dict(slide_cfg.get('transition'))
    effect = transition
    if not cli_overrides.get('transition') and 'effect' in trans_cfg:
        cfg_effect = str(trans_cfg.get('effect'))
        effect = None if cfg_effect == 'none' else cfg_effect
    if not cli_overrides.get('transition_duration'):
        duration = _to_float(trans_cfg.get('duration'), duration)
    if not cli_overrides.get('auto_advance') and 'auto_advance' in trans_cfg:
        auto_advance = _to_float(trans_cfg.get('auto_advance'), auto_advance or 0)
    return effect, duration, auto_advance


def _slide_animation_settings(
    slide_cfg: dict[str, Any],
    animation: str | None,
    duration: float,
    stagger: float,
    trigger: str,
    cli_overrides: dict[str, bool],
) -> tuple[str | None, float, float, str]:
    anim_cfg = _as_dict(slide_cfg.get('animation'))
    effect = animation
    if not cli_overrides.get('animation') and 'effect' in anim_cfg:
        cfg_effect = str(anim_cfg.get('effect'))
        effect = None if cfg_effect == 'none' else cfg_effect
    if not cli_overrides.get('animation_duration'):
        duration = _to_float(anim_cfg.get('duration'), duration)
    if not cli_overrides.get('animation_stagger'):
        stagger = _to_float(anim_cfg.get('stagger'), stagger)
    if not cli_overrides.get('animation_trigger') and anim_cfg.get('trigger'):
        trigger = str(anim_cfg.get('trigger'))
    return effect, duration, stagger, trigger


def _build_sequence_targets(
    anim_targets: list[tuple[int, str]],
    slide_cfg: dict[str, Any],
    animation: str,
    duration: float,
    stagger: float,
    mixed_animation_offset: int,
) -> tuple[list[tuple[int, int, str, float]], int]:
    groups_cfg = _as_dict(slide_cfg.get('groups'))
    ordered: list[tuple[int, int, int, str, dict[str, Any]]] = []
    for idx, (sid, svg_id) in enumerate(anim_targets):
        group_cfg = _as_dict(groups_cfg.get(svg_id))
        if str(group_cfg.get('effect', '')).lower() == 'none':
            continue
        order_value = group_cfg.get('order')
        try:
            order = int(order_value)
            has_order = 0
        except (TypeError, ValueError):
            order = idx
            has_order = 1
        group_entry = dict(group_cfg)
        group_entry['_shape_id'] = sid
        ordered.append((has_order, order, idx, svg_id, group_entry))

    ordered.sort(key=lambda item: (item[0], item[1], item[2]))

    seq_targets: list[tuple[int, int, str, float]] = []
    for seq_idx, (_has_order, _order, _original_idx, _svg_id, group_cfg) in enumerate(ordered):
        shape_id = int(group_cfg['_shape_id'])
        raw_effect = group_cfg.get('effect')
        if raw_effect in ('auto', 'mixed', 'random'):
            effect = pick_animation_effect(
                str(raw_effect), seq_idx, mixed_animation_offset, group_id=_svg_id,
            )
        else:
            effect = str(raw_effect or pick_animation_effect(
                animation, seq_idx, mixed_animation_offset, group_id=_svg_id,
            ))
        item_duration = _to_float(group_cfg.get('duration'), duration)
        delay_seconds = _to_float(
            group_cfg.get('delay'),
            0 if seq_idx == 0 else stagger,
        )
        seq_targets.append((shape_id, int(delay_seconds * 1000), effect, item_duration))

    mixed_count = 0
    if animation == 'mixed':
        mixed_count = sum(1 for _target in seq_targets[1:])
    elif animation == 'auto':
        # 'auto' accumulates a cross-slide offset so the image pool and the
        # unmatched-id fallback rotate as the deck advances. Single-effect
        # semantic matches (title→fade, chart→wipe etc.) are unaffected
        # because they ignore the offset.
        mixed_count = len(seq_targets)
    return seq_targets, mixed_count


def _prerender_legacy_pngs(
    svg_files: list[Path],
    media_dir: Path,
    pixel_width: int,
    pixel_height: int,
    cache_dir: Path | None,
    workers: int,
    verbose: bool,
) -> dict[int, bool]:
    """Render every SVG→PNG into media_dir in parallel.

    Returns {1-based slide index: success}. Falls back to sequential when
    workers<=1 or len(svg_files)<=2.
    """
    results: dict[int, bool] = {}
    targets: list[tuple[int, Path, Path]] = [
        (i, svg, media_dir / f'image{i}.png')
        for i, svg in enumerate(svg_files, 1)
    ]

    if workers <= 1 or len(targets) <= 2:
        for i, svg, png in targets:
            ok = convert_svg_to_png_cached(svg, png, pixel_width, pixel_height, cache_dir)
            results[i] = ok
            if verbose:
                tag = 'cached/ok' if ok else 'failed'
                print(f"  [PNG {i}/{len(targets)}] {svg.name} - {tag}")
        return results

    with ProcessPoolExecutor(max_workers=workers) as pool:
        future_map = {
            pool.submit(
                convert_svg_to_png_cached,
                svg, png, pixel_width, pixel_height, cache_dir,
            ): (i, svg)
            for i, svg, png in targets
        }
        done = 0
        for future in as_completed(future_map):
            i, svg = future_map[future]
            try:
                ok = future.result()
            except Exception as exc:
                ok = False
                if verbose:
                    print(f"  [PNG] {svg.name} - worker error: {exc}")
            results[i] = ok
            done += 1
            if verbose:
                tag = 'cached/ok' if ok else 'failed'
                print(f"  [PNG {done}/{len(targets)}] {svg.name} - {tag}")

    return results


_OPC_UNRESERVED = frozenset(
    "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-._~"
)
_ASCII_LOWER_TRANSLATION = str.maketrans(
    "ABCDEFGHIJKLMNOPQRSTUVWXYZ",
    "abcdefghijklmnopqrstuvwxyz",
)


def _canonical_opc_part_path(path: str) -> str | None:
    """Return an OPC-equivalent package path key, or None when invalid."""
    if (
        not path
        or "\\" in path
        or path.endswith("/")
        or "//" in path
        or any(ord(char) <= 0x20 for char in path)
    ):
        return None
    output: list[str] = []
    index = 0
    while index < len(path):
        char = path[index]
        if char != "%":
            output.append(char)
            index += 1
            continue
        if index + 2 >= len(path) or not re.fullmatch(r"[0-9A-Fa-f]{2}", path[index + 1:index + 3]):
            return None
        value = int(path[index + 1:index + 3], 16)
        decoded = chr(value)
        if value in {0, ord("/"), ord("\\")}:
            return None
        output.append(decoded if decoded in _OPC_UNRESERVED else f"%{value:02X}")
        index += 3

    decoded_path = "".join(output)
    if decoded_path.rsplit("/", 1)[-1] in {".", ".."}:
        return None
    normalized = posixpath.normpath(decoded_path)
    if (
        not normalized
        or normalized in {".", ".."}
        or normalized.startswith("/")
        or normalized.startswith("../")
    ):
        return None
    return normalized.translate(_ASCII_LOWER_TRANSLATION)


def _source_part_for_rels(rels_path: str) -> str | None:
    """Return the source part path represented by a relationship sidecar."""
    filename = posixpath.basename(rels_path)
    if filename == ".rels" or not filename.endswith(".rels"):
        return None
    source_dir = posixpath.dirname(posixpath.dirname(rels_path))
    source_name = filename[:-len(".rels")]
    return posixpath.join(source_dir, source_name) if source_dir else source_name


def _resolve_internal_opc_target(rels_path: str, target: str) -> str | None:
    """Resolve one valid internal OPC Target to its canonical package key."""
    target_path_query = target.split("#", 1)[0]
    if (
        "\\" in target
        or "?" in target_path_query
        or any(ord(char) <= 0x20 for char in target)
    ):
        return None
    try:
        parsed = urlsplit(target)
    except ValueError:
        return None
    if parsed.scheme or parsed.netloc or parsed.query:
        return None

    source_part = _source_part_for_rels(rels_path)
    if parsed.path.startswith("/"):
        resolved = parsed.path[1:]
    elif parsed.path:
        base_dir = posixpath.dirname(source_part) if source_part else ""
        resolved = posixpath.join(base_dir, parsed.path) if base_dir else parsed.path
    elif source_part and "#" in target:
        resolved = source_part
    else:
        return None
    return _canonical_opc_part_path(resolved)


def _verify_internal_rels_targets(extract_dir: Path) -> list[str]:
    """Return a list of dangling internal Targets across every .rels in the package.

    Each entry is formatted as "<rels-path> -> <missing-target>". An empty list
    means every internal Target resolves to a real file in the package.
    """
    package_parts: set[str] = set()
    for path in extract_dir.rglob("*"):
        if not path.is_file():
            continue
        key = _canonical_opc_part_path(path.relative_to(extract_dir).as_posix())
        if key is not None:
            package_parts.add(key)
    problems: list[str] = []
    for rels_path in extract_dir.rglob('*.rels'):
        rels_rel = rels_path.relative_to(extract_dir).as_posix()
        try:
            root = ET.parse(rels_path).getroot()
        except ET.ParseError as exc:
            problems.append(f'{rels_rel} -> <invalid relationships XML: {exc}>')
            continue
        for elem in root:
            attrs = _relationship_attrs(elem)
            if attrs.get('TargetMode', '').lower() == 'external':
                continue
            target = attrs.get('Target')
            if not target:
                problems.append(f'{rels_rel} -> <missing Target>')
                continue
            resolved = _resolve_internal_opc_target(rels_rel, target)
            if resolved is None:
                problems.append(f'{rels_rel} -> <invalid Target {target!r}>')
                continue
            if resolved not in package_parts:
                problems.append(f'{rels_rel} -> {resolved}')
    return problems


def _presentation_format(width: float, height: float) -> str:
    """Map the slide aspect ratio to PowerPoint's PresentationFormat label.
    Non-standard ratios (square, portrait, banner crops) report 'Custom'.
    """
    if width <= 0 or height <= 0:
        return 'Custom'
    ratio = width / height
    for target, label in (
        (4 / 3, 'On-screen Show (4:3)'),
        (16 / 9, 'On-screen Show (16:9)'),
        (16 / 10, 'On-screen Show (16:10)'),
    ):
        if abs(ratio - target) < 0.02:
            return label
    return 'Custom'


def _stamp_docprops(
    extract_dir: Path,
    slide_count: int,
    pres_format: str,
    meta: dict[str, Any] | None = None,
) -> None:
    """Overwrite the misleading python-pptx default metadata with accurate
    values. Factual fields (slide count, export timestamp, presentation format,
    application) are always machine-derived. Authored fields — including the
    title — come solely from an optional per-project ``metadata.json``
    (``meta``); whatever it omits stays blank. ``lastModifiedBy`` follows
    ``creator`` rather than ever carrying the base template's author or a tool
    name. No field is guessed from slide content: a blank title is preferable
    to an unreliable heuristic pick.
    """
    meta = meta or {}

    def field(key: str, default: str = '') -> str:
        value = meta.get(key)
        return value.strip() if isinstance(value, str) and value.strip() else default

    title = field('title')
    creator = field('creator')

    now = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')

    core_path = extract_dir / 'docProps' / 'core.xml'
    if core_path.exists():
        core_path.write_text(
            "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n"
            '<cp:coreProperties '
            'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" '
            'xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{escape(title)}</dc:title>'
            f'<dc:subject>{escape(field("subject"))}</dc:subject>'
            f'<dc:creator>{escape(creator)}</dc:creator>'
            f'<cp:keywords>{escape(field("keywords"))}</cp:keywords>'
            f'<dc:description>{escape(field("description"))}</dc:description>'
            f'<dc:language>{escape(field("language"))}</dc:language>'
            f'<cp:lastModifiedBy>{escape(creator)}</cp:lastModifiedBy>'
            '<cp:revision>1</cp:revision>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>'
            f'<cp:category>{escape(field("category"))}</cp:category>'
            f'<cp:contentStatus>{escape(field("contentStatus"))}</cp:contentStatus>'
            '</cp:coreProperties>',
            encoding='utf-8',
        )

    app_path = extract_dir / 'docProps' / 'app.xml'
    if app_path.exists():
        app = app_path.read_text(encoding='utf-8')
        app = re.sub(r'<Slides>.*?</Slides>', f'<Slides>{slide_count}</Slides>', app)
        app = re.sub(
            r'<Company>.*?</Company>',
            f'<Company>{escape(field("company"))}</Company>',
            app,
        )
        app = re.sub(
            r'<Manager>.*?</Manager>',
            f'<Manager>{escape(field("manager"))}</Manager>',
            app,
        )
        app = re.sub(
            r'<Application>.*?</Application>',
            '<Application>Microsoft Office PowerPoint</Application>',
            app,
        )
        app = re.sub(
            r'<PresentationFormat>.*?</PresentationFormat>',
            f'<PresentationFormat>{escape(pres_format)}</PresentationFormat>',
            app,
        )
        app_path.write_text(app, encoding='utf-8')


def create_pptx_with_native_svg(
    svg_files: list[Path],
    output_path: Path,
    canvas_format: str | None = None,
    verbose: bool = True,
    transition: str | None = 'fade',
    transition_duration: float = 0.5,
    auto_advance: float | None = None,
    use_compat_mode: bool = True,
    notes: dict[str, str] | None = None,
    enable_notes: bool = True,
    use_native_shapes: bool = False,
    animation: str | None = None,
    animation_duration: float = 0.4,
    animation_stagger: float = 0.5,
    animation_trigger: str = 'after-previous',
    animation_config: dict[str, Any] | None = None,
    animation_cli_overrides: dict[str, bool] | None = None,
    narration_audio: dict[str, Path] | None = None,
    use_narration_timings: bool = False,
    narration_padding: float = 0.5,
    absolute_link_base: Path | None = None,
    cache_dir: Path | None = None,
    workers: int | None = None,
    merge_paragraphs: bool = True,
    image_optimize: bool = True,
    image_max_dimension: int | None = 2560,
    image_sizing: str = 'cap',
    image_scale: float = 2.0,
    image_quality: int = 85,
    native_objects: bool = False,
    conversion_trace_path: Path | None = None,
    doc_metadata: dict[str, Any] | None = None,
    pptx_structure: str = "baseline",
) -> bool:
    """Create a PPTX file with native SVG.

    Args:
        svg_files: List of SVG files.
        output_path: Output PPTX path.
        canvas_format: Canvas format key.
        verbose: Whether to output detailed information.
        transition: Transition effect name.
        transition_duration: Transition duration in seconds.
        auto_advance: Auto-advance interval in seconds.
        use_compat_mode: Use Office compatibility mode (PNG + SVG dual format).
        notes: Notes dict, key is SVG stem, value is notes content.
        enable_notes: Whether to enable notes embedding.
        use_native_shapes: Convert SVG to native DrawingML shapes.
        animation: Per-element entrance animation mode (single effect name,
            'mixed', 'random', or None to disable). Native shapes mode only.
        animation_duration: Per-element entrance duration in seconds.
        animation_stagger: Delay between elements in ``after-previous``
            trigger mode (seconds). Ignored otherwise.
        animation_trigger: PowerPoint Start mode — ``'after-previous'`` (default),
            ``'on-click'``, or ``'with-previous'``.
        animation_config: Optional sidecar overrides loaded from animations.json.
        animation_cli_overrides: Flags indicating explicit CLI overrides.
        narration_audio: Optional dict mapping SVG stem to narration audio file.
        use_narration_timings: Whether to set slide auto-advance from audio duration.
        narration_padding: Extra seconds added after each narration before advancing.
        image_optimize: Whether native export downscales oversized raster images.
        image_max_dimension: Maximum optimized image dimension in pixels.
        image_sizing: ``cap`` only limits source dimensions; ``display`` sizes
            from rendered SVG boxes.
        image_scale: Target image pixels per SVG display pixel.
        image_quality: JPEG quality used for opaque optimized rasters.
        native_objects: Convert explicit ``data-pptx-native`` table/chart
            markers to native PowerPoint objects. Default off.
        conversion_trace_path: Optional JSON path for native conversion diagnostics.
        pptx_structure: PPTX structure strategy. ``baseline`` promotes safe
            shared native backgrounds and leading chrome to slide masters;
            ``flat`` keeps generated backgrounds/chrome slide-local.

    Returns:
        Whether all slides were successfully created.
    """
    if not svg_files:
        print("Error: No SVG files found")
        return False

    # Native shapes mode takes priority over compat mode
    if use_native_shapes:
        use_compat_mode = False
    if pptx_structure not in {"baseline", "flat"}:
        raise ValueError(f"Unsupported pptx_structure: {pptx_structure}")

    # Check compatibility mode dependencies
    renderer_name, renderer_status, renderer_hint = get_png_renderer_info()
    if not use_native_shapes and use_compat_mode and PNG_RENDERER is None:
        print("Warning: No PNG rendering library installed, cannot use compatibility mode")
        print(f"  {renderer_hint}")
        print("  Will use pure SVG mode (may not display in Office LTSC 2021 and similar versions)")
        use_compat_mode = False

    # Auto-detect canvas format or get dimensions from viewBox
    custom_pixels: tuple[int, int] | None = None
    if canvas_format is None:
        canvas_format = detect_format_from_svg(svg_files[0])
        if canvas_format and verbose:
            format_name = CANVAS_FORMATS.get(canvas_format, {}).get('name', canvas_format)
            print(f"  Detected canvas format: {format_name}")

    if canvas_format is None:
        custom_pixels = get_viewbox_dimensions(svg_files[0])
        if custom_pixels and verbose:
            print(f"  Using SVG viewBox dimensions: {custom_pixels[0]} x {custom_pixels[1]} px")

    if canvas_format is None and custom_pixels is None:
        canvas_format = 'ppt169'
        if verbose:
            print(f"  Using default format: PPT 16:9")

    width_emu, height_emu = get_slide_dimensions(canvas_format or 'ppt169', custom_pixels)
    pixel_width, pixel_height = get_pixel_dimensions(canvas_format or 'ppt169', custom_pixels)

    if verbose:
        print(f"  Slide dimensions: {pixel_width} x {pixel_height} px")
        print(f"  SVG file count: {len(svg_files)}")
        if use_native_shapes:
            print(f"  Mode: Native DrawingML shapes (directly editable)")
            print(
                "  Native table/chart objects: "
                f"{'Enabled' if native_objects else 'Disabled'}"
            )
            print(f"  PPTX structure: {pptx_structure}")
            if image_optimize:
                if image_sizing == 'display':
                    image_mode = (
                        f"display scale {image_scale:g}, "
                        f"max {image_max_dimension or 'unlimited'} px"
                    )
                else:
                    image_mode = f"cap max {image_max_dimension or 'unlimited'} px"
                print(
                    "  Image optimization: Enabled "
                    f"({image_mode}, JPEG q{image_quality})"
                )
            else:
                print("  Image optimization: Disabled")
        elif use_compat_mode:
            print(f"  Compatibility mode: Enabled (PNG + SVG dual format)")
            print(f"  PNG renderer: {renderer_name} {renderer_status}")
        else:
            print(f"  Compatibility mode: Disabled (pure SVG)")
        if transition:
            trans_name = TRANSITIONS.get(transition, {}).get('name', transition) if TRANSITIONS else transition
            print(f"  Transition effect: {trans_name}")
        if enable_notes and notes:
            print(f"  Speaker notes: {len(notes)} page(s)")
        elif enable_notes:
            print(f"  Speaker notes: Enabled (no notes files found)")
        else:
            print(f"  Speaker notes: Disabled")
        print()

    animation_cli_overrides = animation_cli_overrides or {}

    temp_dir = _create_writable_work_dir(output_path)

    try:
        # Create base PPTX with python-pptx
        prs = Presentation()
        prs.slide_width = width_emu
        prs.slide_height = height_emu

        blank_layout = prs.slide_layouts[6]
        for _ in svg_files:
            prs.slides.add_slide(blank_layout)

        base_pptx = temp_dir / 'base.pptx'
        prs.save(str(base_pptx))

        # Extract PPTX
        extract_dir = temp_dir / 'pptx_content'
        with zipfile.ZipFile(base_pptx, 'r') as zf:
            zf.extractall(extract_dir)
        structure = _read_slide_layout_targets(extract_dir, len(svg_files))

        media_dir = extract_dir / 'ppt' / 'media'
        media_dir.mkdir(exist_ok=True)

        prerender_results: dict[int, bool] | None = None
        if not use_native_shapes and use_compat_mode and PNG_RENDERER is not None:
            if workers is None:
                resolved_workers = min(os.cpu_count() or 2, len(svg_files), 8)
            else:
                resolved_workers = max(0, workers)
            if verbose:
                cache_label = str(cache_dir) if cache_dir else 'disabled'
                mode = f'parallel x{resolved_workers}' if resolved_workers > 1 else 'sequential'
                print(f"  Pre-rendering PNGs ({mode}, cache: {cache_label})")
            prerender_results = _prerender_legacy_pngs(
                svg_files, media_dir, pixel_width, pixel_height,
                cache_dir, resolved_workers, verbose,
            )
            if verbose:
                print()

        success_count = 0
        has_any_image = False
        media_cache: dict[tuple[str, str], str] = {}
        image_exts_used: set[str] = set()
        package_exts_used: set[str] = set()
        package_content_overrides: dict[str, str] = {}
        notes_slides_created: set[int] = set()
        narration_slides_created: set[int] = set()
        audio_exts_used: set[str] = set()
        mixed_animation_offset = 0
        conversion_trace: list[dict[str, Any]] | None = [] if conversion_trace_path else None
        structure_trace: list[dict[str, Any]] | None = (
            [] if use_native_shapes and pptx_structure == "baseline" else None
        )

        for i, svg_path in enumerate(svg_files, 1):
            slide_num = i

            try:
                # ---- Native shapes mode ----
                if use_native_shapes:
                    slide_cfg = _slide_config(animation_config, svg_path.stem)
                    (
                        slide_xml,
                        media_files_dict,
                        rel_entries,
                        anim_targets,
                        package_files_dict,
                        content_type_overrides,
                    ) = (
                        convert_svg_to_slide_shapes(
                            svg_path, slide_num=slide_num, verbose=verbose,
                            merge_paragraphs=merge_paragraphs,
                            image_optimize=image_optimize,
                            image_max_dimension=image_max_dimension,
                            image_sizing=image_sizing,
                            image_scale=image_scale,
                            image_quality=image_quality,
                            native_objects=native_objects,
                            trace_out=conversion_trace
                            if conversion_trace is not None
                            else structure_trace,
                        )
                    )
                    slide_transition, slide_transition_duration, slide_auto_advance = (
                        _slide_transition_settings(
                            slide_cfg,
                            transition,
                            transition_duration,
                            auto_advance,
                            animation_cli_overrides,
                        )
                    )
                    (
                        slide_animation,
                        slide_animation_duration,
                        slide_animation_stagger,
                        slide_animation_trigger,
                    ) = _slide_animation_settings(
                        slide_cfg,
                        animation,
                        animation_duration,
                        animation_stagger,
                        animation_trigger,
                        animation_cli_overrides,
                    )

                    # Inject SVG hyperlinks as transparent overlay shapes
                    # (must precede transition/timing injection so the spTree
                    # is finalized before <p:transition>/<p:timing> are appended).
                    raw_links = extract_links(svg_path, width_emu, height_emu, absolute_link_base=absolute_link_base)
                    if raw_links:
                        used_rids = [
                            int(re.match(r'rId(\d+)', r['id']).group(1))
                            for r in rel_entries
                            if re.match(r'rId(\d+)', r['id'])
                        ]
                        next_rid = max(used_rids, default=1) + 1
                        used_shape_ids = [
                            int(m.group(1))
                            for m in re.finditer(r'<p:cNvPr id="(\d+)"', slide_xml)
                        ]
                        next_shape_id = max(used_shape_ids, default=1) + 1
                        link_shapes = []
                        for lk in raw_links:
                            rid = f'rId{next_rid}'
                            next_rid += 1
                            rel_entries.append({
                                'id': rid,
                                'type': _HYPERLINK_REL_TYPE,
                                'target': lk['href'],
                                'target_mode': 'External',
                            })
                            link_shapes.append(link_shape_xml(
                                shape_id=next_shape_id,
                                href_rid=rid,
                                x=lk['x'], y=lk['y'],
                                w=lk['w'], h=lk['h'],
                            ))
                            next_shape_id += 1
                        if link_shapes:
                            slide_xml = slide_xml.replace(
                                '</p:spTree>',
                                '\n' + '\n'.join(link_shapes) + '\n      </p:spTree>',
                            )

                    # Order matters: OOXML schema requires <p:transition>
                    # to precede <p:timing> inside <p:sld>. Both use the same
                    # </p:sld> string-replace anchor, so transition must be
                    # injected first and timing second.
                    if slide_transition and ANIMATIONS_AVAILABLE and create_transition_xml:
                        transition_xml = '\n' + create_transition_xml(
                            effect=slide_transition,
                            duration=slide_transition_duration,
                            advance_after=slide_auto_advance,
                        )
                        slide_xml = slide_xml.replace(
                            '</p:sld>',
                            transition_xml + '\n</p:sld>',
                        )

                    if (slide_animation and slide_animation != 'none'
                            and create_sequence_timing_xml
                            and pick_animation_effect
                            and anim_targets):
                        seq_targets, mixed_count = _build_sequence_targets(
                            anim_targets,
                            slide_cfg,
                            slide_animation,
                            slide_animation_duration,
                            slide_animation_stagger,
                            mixed_animation_offset,
                        )
                        if slide_animation in ('mixed', 'auto'):
                            mixed_animation_offset += mixed_count
                        timing_xml = '\n' + create_sequence_timing_xml(
                            seq_targets, duration=slide_animation_duration,
                            trigger=slide_animation_trigger,
                        )
                        slide_xml = slide_xml.replace(
                            '</p:sld>',
                            timing_xml + '\n</p:sld>',
                        )

                    # Write slide XML
                    slide_xml_path = extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                    with open(slide_xml_path, 'w', encoding='utf-8') as f:
                        f.write(slide_xml)

                    # Write media files
                    media_name_map: dict[str, str] = {}
                    for media_name, media_data in media_files_dict.items():
                        ext = media_name.rsplit('.', 1)[-1].lower()
                        media_hash = hashlib.sha256(media_data).hexdigest()
                        cache_key = (ext, media_hash)
                        cached_name = media_cache.get(cache_key)

                        if cached_name is None:
                            cached_name = f'image_{media_hash[:16]}.{ext}'
                            media_cache[cache_key] = cached_name
                            with open(media_dir / cached_name, 'wb') as f:
                                f.write(media_data)

                        media_name_map[media_name] = cached_name

                    for rel in rel_entries:
                        target = rel.get('target', '')
                        if not target.startswith('../media/'):
                            continue
                        media_name = target.split('../media/', 1)[1]
                        mapped_name = media_name_map.get(media_name)
                        if mapped_name:
                            rel['target'] = f'../media/{mapped_name}'

                    # Write non-media OOXML package parts produced by native
                    # object converters, e.g. chart XML, chart rels, and
                    # embedded workbooks.
                    for part_name, part_data in package_files_dict.items():
                        package_path = extract_dir / part_name
                        package_path.parent.mkdir(parents=True, exist_ok=True)
                        with open(package_path, 'wb') as f:
                            f.write(part_data)
                        suffix = package_path.suffix.lstrip('.').lower()
                        if suffix:
                            package_exts_used.add(suffix)
                    package_content_overrides.update(content_type_overrides)

                    # Build relationships XML
                    rels_dir = extract_dir / 'ppt' / 'slides' / '_rels'
                    rels_dir.mkdir(exist_ok=True)
                    rels_path = rels_dir / f'slide{slide_num}.xml.rels'

                    extra_rels = ''
                    for rel in rel_entries:
                        mode_attr = (
                            f' TargetMode="{rel["target_mode"]}"'
                            if rel.get('target_mode') else ''
                        )
                        extra_rels += (
                            f'\n  <Relationship Id="{rel["id"]}" '
                            f'Type="{rel["type"]}" Target="{rel["target"]}"{mode_attr}/>'
                        )

                    rels_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
                Type="{SLIDE_LAYOUT_REL_TYPE}"
                Target="{structure.slide_layout_target(slide_num)}"/>{extra_rels}
</Relationships>'''
                    with open(rels_path, 'w', encoding='utf-8') as f:
                        f.write(rels_xml)

                    # Track image formats for Content_Types
                    for media_name in media_name_map.values():
                        ext = media_name.rsplit('.', 1)[-1].lower()
                        _content_type_for_extension(ext)
                        image_exts_used.add(ext)
                        has_any_image = True

                # ---- Legacy SVG embedding mode ----
                else:
                    slide_cfg = _slide_config(animation_config, svg_path.stem)
                    slide_transition, slide_transition_duration, slide_auto_advance = (
                        _slide_transition_settings(
                            slide_cfg,
                            transition,
                            transition_duration,
                            auto_advance,
                            animation_cli_overrides,
                        )
                    )
                    svg_filename = f'image{i}.svg'
                    png_filename = f'image{i}.png'
                    png_rid = 'rId2'
                    svg_rid = 'rId3' if use_compat_mode else 'rId2'

                    shutil.copy(svg_path, media_dir / svg_filename)

                    slide_has_png = False
                    if use_compat_mode:
                        if prerender_results is not None:
                            png_success = prerender_results.get(i, False)
                        else:
                            png_path = media_dir / png_filename
                            png_success = convert_svg_to_png(
                                svg_path, png_path,
                                width=pixel_width, height=pixel_height,
                            )
                        if png_success:
                            slide_has_png = True
                            has_any_image = True
                            image_exts_used.add('png')
                        else:
                            if verbose:
                                print(
                                    f"  [{i}/{len(svg_files)}] {svg_path.name} - "
                                    "PNG generation failed, using pure SVG"
                                )
                            svg_rid = 'rId2'

                    # Extract SVG hyperlinks and assign rIds (starting after rId3)
                    raw_links = extract_links(svg_path, width_emu, height_emu, absolute_link_base=absolute_link_base)
                    link_rels = []
                    link_regions = []
                    for li, lk in enumerate(raw_links):
                        rid = f'rId{4 + li}'
                        link_rels.append({'rid': rid, 'href': lk['href']})
                        link_regions.append({
                            'href_rid': rid,
                            'x': lk['x'], 'y': lk['y'],
                            'w': lk['w'], 'h': lk['h'],
                        })

                    slide_xml_path = extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                    slide_xml = create_slide_xml_with_svg(
                        slide_num,
                        png_rid=png_rid, svg_rid=svg_rid,
                        width_emu=width_emu, height_emu=height_emu,
                        transition=slide_transition,
                        transition_duration=slide_transition_duration,
                        auto_advance=slide_auto_advance,
                        use_compat_mode=(use_compat_mode and slide_has_png),
                        link_regions=link_regions or None,
                    )
                    with open(slide_xml_path, 'w', encoding='utf-8') as f:
                        f.write(slide_xml)

                    rels_dir = extract_dir / 'ppt' / 'slides' / '_rels'
                    rels_dir.mkdir(exist_ok=True)
                    rels_path = rels_dir / f'slide{slide_num}.xml.rels'
                    rels_xml = create_slide_rels_xml(
                        png_rid=png_rid, png_filename=png_filename,
                        svg_rid=svg_rid, svg_filename=svg_filename,
                        use_compat_mode=(use_compat_mode and slide_has_png),
                        link_rels=link_rels or None,
                        slide_layout_target=structure.slide_layout_target(slide_num),
                    )
                    with open(rels_path, 'w', encoding='utf-8') as f:
                        f.write(rels_xml)

                # --- Process notes (shared between native and legacy mode) ---
                notes_content = ''
                if enable_notes:
                    svg_stem = svg_path.stem
                    notes_content = notes.get(svg_stem, '') if notes else ''
                    notes_text = markdown_to_plain_text(notes_content) if notes_content else ''
                    if notes_text:
                        _ensure_notes_master(extract_dir)

                        notes_slides_dir = extract_dir / 'ppt' / 'notesSlides'
                        notes_slides_dir.mkdir(exist_ok=True)

                        notes_xml_path = notes_slides_dir / f'notesSlide{slide_num}.xml'
                        notes_xml = create_notes_slide_xml(slide_num, notes_text)
                        with open(notes_xml_path, 'w', encoding='utf-8') as f:
                            f.write(notes_xml)

                        notes_rels_dir = notes_slides_dir / '_rels'
                        notes_rels_dir.mkdir(exist_ok=True)
                        notes_rels_path = notes_rels_dir / f'notesSlide{slide_num}.xml.rels'
                        notes_rels_xml = create_notes_slide_rels_xml(slide_num)
                        with open(notes_rels_path, 'w', encoding='utf-8') as f:
                            f.write(notes_rels_xml)

                        _append_relationship(
                            rels_path,
                            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide',
                            f'../notesSlides/notesSlide{slide_num}.xml',
                        )
                        notes_slides_created.add(slide_num)

                # --- Process narration audio (shared between native and legacy mode) ---
                svg_stem = svg_path.stem
                audio_path = narration_audio.get(svg_stem) if narration_audio else None
                if audio_path:
                    slide_xml_path = extract_dir / 'ppt' / 'slides' / f'slide{slide_num}.xml'
                    rels_path = extract_dir / 'ppt' / 'slides' / '_rels' / f'slide{slide_num}.xml.rels'

                    ext = audio_path.suffix.lower()
                    media_name = f'narration{slide_num}{ext}'
                    shutil.copy2(audio_path, media_dir / media_name)
                    audio_exts_used.add(ext)

                    poster_name = 'narration_poster.png'
                    poster_path = media_dir / poster_name
                    if not poster_path.exists():
                        poster_path.write_bytes(AUDIO_MARKER_PNG_BYTES)
                    has_any_image = True
                    image_exts_used.add('png')

                    media_rid = _append_relationship(
                        rels_path,
                        MEDIA_REL_TYPE,
                        f'../media/{media_name}',
                    )
                    audio_rid = _append_relationship(
                        rels_path,
                        AUDIO_REL_TYPE,
                        f'../media/{media_name}',
                    )
                    poster_rid = _append_relationship(
                        rels_path,
                        IMAGE_REL_TYPE,
                        f'../media/{poster_name}',
                    )

                    slide_xml = slide_xml_path.read_text(encoding='utf-8')
                    narration_shape_id = next_shape_id(slide_xml)
                    slide_xml = inject_narration(
                        slide_xml,
                        shape_id=narration_shape_id,
                        shape_name=media_name,
                        audio_rid=audio_rid,
                        media_rid=media_rid,
                        poster_rid=poster_rid,
                    )

                    if use_narration_timings:
                        duration = probe_audio_duration(audio_path)
                        if duration is None:
                            raise RuntimeError(
                                f"Unable to read narration duration with ffprobe: {audio_path}"
                            )
                        slide_xml = apply_recorded_timing(
                            slide_xml,
                            advance_after=duration + narration_padding,
                            transition_duration=slide_transition_duration,
                            transition_effect=slide_transition or 'fade',
                        )
                    slide_xml_path.write_text(slide_xml, encoding='utf-8')
                    narration_slides_created.add(slide_num)

                if verbose:
                    if use_native_shapes:
                        mode_str = " (Native)"
                    elif use_compat_mode and not use_native_shapes:
                        mode_str = " (PNG+SVG)" if has_any_image else " (SVG)"
                    else:
                        mode_str = " (SVG)"
                    has_notes = slide_num in notes_slides_created
                    notes_str = " +notes" if has_notes else ""
                    narration_str = " +narration" if slide_num in narration_slides_created else ""
                    print(f"  [{i}/{len(svg_files)}] {svg_path.name}{mode_str}{notes_str}{narration_str}")

                success_count += 1

            except Exception as e:
                if verbose:
                    print(f"  [{i}/{len(svg_files)}] {svg_path.name} - Error: {e}")
                if use_native_shapes:
                    raise

        if (
            use_native_shapes
            and pptx_structure == "baseline"
            and success_count == len(svg_files)
        ):
            _promote_common_slide_backgrounds_to_masters(
                extract_dir,
                structure,
                len(svg_files),
                verbose=verbose,
            )
            _promote_common_chrome_shapes_to_masters(
                extract_dir,
                structure,
                len(svg_files),
                conversion_trace if conversion_trace is not None else structure_trace,
                verbose=verbose,
            )

        # Update [Content_Types].xml
        content_types_path = extract_dir / '[Content_Types].xml'
        with open(content_types_path, 'r', encoding='utf-8') as f:
            content_types = f.read()

        if not use_native_shapes:
            content_types = _add_default_content_type(content_types, 'svg', 'image/svg+xml')
        for ext in sorted(image_exts_used):
            content_types = _add_default_content_type(
                content_types,
                ext,
                _content_type_for_extension(ext),
            )
        if 'xlsx' in package_exts_used:
            content_types = _add_default_content_type(
                content_types,
                'xlsx',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
        for part_name, content_type in sorted(package_content_overrides.items()):
            content_types = _add_content_type_override(content_types, part_name, content_type)
        with open(content_types_path, 'w', encoding='utf-8') as f:
            f.write(content_types)

        if audio_exts_used:
            for ext in sorted(audio_exts_used):
                content_type = AUDIO_CONTENT_TYPES.get(ext)
                if content_type:
                    content_types = _add_default_content_type(content_types, ext, content_type)
            if 'Extension="png"' not in content_types:
                content_types = _add_default_content_type(content_types, 'png', 'image/png')
            with open(content_types_path, 'w', encoding='utf-8') as f:
                f.write(content_types)

        # Add notes master / slides content types
        if enable_notes and notes_slides_created:
            notes_theme_override = (
                '  <Override PartName="/ppt/theme/theme2.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
            )
            if notes_theme_override not in content_types:
                content_types = content_types.replace(
                    '</Types>',
                    notes_theme_override + '\n</Types>',
                )
            notes_master_override = (
                '  <Override PartName="/ppt/notesMasters/notesMaster1.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml"/>'
            )
            if notes_master_override not in content_types:
                content_types = content_types.replace(
                    '</Types>',
                    notes_master_override + '\n</Types>',
                )
            for i in sorted(notes_slides_created):
                override = (
                    f'  <Override PartName="/ppt/notesSlides/notesSlide{i}.xml" '
                    f'ContentType="application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml"/>'
                )
                if override not in content_types:
                    content_types = content_types.replace('</Types>', override + '\n</Types>')
            with open(content_types_path, 'w', encoding='utf-8') as f:
                f.write(content_types)

        rels_problems = _verify_internal_rels_targets(extract_dir)
        if rels_problems:
            details = '\n'.join(f'  - {p}' for p in rels_problems)
            raise RuntimeError(
                'PPTX package contains dangling internal relationship targets; '
                'PowerPoint will report the file as corrupt:\n' + details
            )

        # Replace the python-pptx base-template metadata (stale "Steve Canny"
        # author, 2013 dates, "generated using python-pptx", Slides=0) with
        # accurate, tool-neutral document properties.
        pres_format = _presentation_format(width_emu, height_emu)
        _stamp_docprops(extract_dir, len(svg_files), pres_format, doc_metadata)

        # Repackage PPTX to a temporary file first. The public output path is
        # replaced only after every slide and relationship has succeeded.
        temp_output_path = temp_dir / 'result.pptx'
        with zipfile.ZipFile(temp_output_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            for file_path in extract_dir.rglob('*'):
                if file_path.is_file():
                    arcname = file_path.relative_to(extract_dir)
                    zf.write(file_path, arcname)
        shutil.move(str(temp_output_path), str(output_path))
        permission_warnings = _relax_output_permissions(output_path)

        if conversion_trace_path and conversion_trace is not None:
            conversion_trace_path.parent.mkdir(parents=True, exist_ok=True)
            payload = {
                'output': str(output_path),
                'slide_count': len(svg_files),
                'slides': conversion_trace,
            }
            conversion_trace_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding='utf-8',
            )

        if verbose:
            print()
            print(f"[Done] Saved: {output_path}")
            for warning in permission_warnings:
                print(f"  [warn] {warning}")
            if conversion_trace_path and conversion_trace is not None:
                print(f"  Trace: {conversion_trace_path}")
            print(f"  Succeeded: {success_count}, Failed: {len(svg_files) - success_count}")
            if use_compat_mode and has_any_image:
                print(f"  Mode: Office compatibility mode (supports all Office versions)")
                if PNG_RENDERER == 'svglib' and renderer_hint:
                    print(f"  [Tip] {renderer_hint}")

        return success_count == len(svg_files)

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)
